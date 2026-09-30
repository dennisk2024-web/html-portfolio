#!/usr/bin/env python3
"""Generate Dennis Kariuki's GOD TV Partnership Co-ordinator cover letter."""

from datetime import date

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt, RGBColor


NAVY = RGBColor(0x1A, 0x36, 0x5D)
BODY = RGBColor(0x2D, 0x2D, 0x2D)
MUTED = RGBColor(0x4A, 0x55, 0x68)


def set_run_font(run, name="Calibri", size=11, bold=False, italic=False, color=BODY):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def add_bottom_border(paragraph, color="C4A44A", size="12"):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "8")
    bottom.set(qn("w:color"), color)
    pBdr.append(bottom)
    pPr.append(pBdr)


def add_paragraph(
    doc,
    text,
    *,
    size=11,
    bold=False,
    italic=False,
    color=BODY,
    align=WD_ALIGN_PARAGRAPH.LEFT,
    space_before=0,
    space_after=10,
    line_spacing=1.15,
):
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = line_spacing
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic, color=color)
    return p


def build():
    doc = Document()

    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    name = add_paragraph(
        doc,
        "Dennis Kariuki",
        size=22,
        bold=True,
        color=NAVY,
        space_after=2,
        line_spacing=1.0,
    )
    add_bottom_border(name)

    add_paragraph(
        doc,
        "United Kingdom  ·  dennisk2024@gmail.com",
        size=11,
        color=MUTED,
        space_before=6,
        space_after=12,
        line_spacing=1.0,
    )

    add_paragraph(
        doc,
        date(2026, 9, 30).strftime("%-d %B %Y"),
        size=11,
        color=BODY,
        space_after=10,
        line_spacing=1.0,
    )

    add_paragraph(doc, "Personnel Team", size=11, bold=True, color=NAVY, space_after=0, line_spacing=1.0)
    add_paragraph(doc, "GOD TV", size=11, color=BODY, space_after=0, line_spacing=1.0)
    add_paragraph(doc, "personnel@god.tv", size=11, color=BODY, space_after=12, line_spacing=1.0)

    add_paragraph(
        doc,
        "Application: Partnership Co-ordinator (UK Partnerships Team)",
        size=11,
        bold=True,
        italic=True,
        color=NAVY,
        space_after=12,
        line_spacing=1.0,
    )

    add_paragraph(doc, "Dear Personnel Team,", size=11, space_after=8)

    body_kwargs = {
        "align": WD_ALIGN_PARAGRAPH.JUSTIFY,
        "space_after": 8,
        "line_spacing": 1.12,
    }

    add_paragraph(
        doc,
        "I am writing to apply for the Partnership Co-ordinator role with GOD TV. "
        "I am a practising and committed Christian, confident and willing to pray with "
        "supporters, and I want to serve viewers across the UK, Europe, the Nordics "
        "and the US from a place of genuine faith rather than a sales mindset. This "
        "role stands out because it is front-line ministry: encouraging partners, "
        "sharing testimonies, and creating God-honouring opportunities for people to "
        "give into the work of the Gospel.",
        **body_kwargs,
    )

    add_paragraph(
        doc,
        "I serve in my church’s media department, where I film Sunday services, edit "
        "videos and look after social media. That work has taught me how to tell "
        "stories that help people encounter God, serve a live congregation with care, "
        "and stay present with people online after the service ends. I understand "
        "audience-led ministry from the inside: capturing what God is doing, shaping "
        "it with integrity, and sharing it in a way that invites a response. This sits "
        "close to GOD TV’s calling as a global Christian media network, and it is the "
        "same heart I would bring to partnership conversations, thank-you calls, "
        "first-time donor follow-up and social media engagement.",
        **body_kwargs,
    )

    add_paragraph(
        doc,
        "Through filming, editing and social media I have developed the skills this "
        "role depends on: building rapport quickly, listening well, and speaking with "
        "warmth and clarity. I regularly respond to messages, encourage people in "
        "their walk with Christ, and help the church stay connected beyond Sunday. I "
        "am comfortable on the phone, WhatsApp and social platforms, and I take "
        "administration seriously — keeping details accurate, following through on "
        "requests, and supporting the wider team when needed. I do not treat "
        "fundraising as persuasion for its own sake. I see it as stewardship: helping "
        "people find the right opportunity to partner, then caring for that "
        "relationship over time.",
        **body_kwargs,
    )

    add_paragraph(
        doc,
        "I work well towards clear goals and manage my own time, which matters in a "
        "remote telephone-based role with an annual income target. I also think "
        "practically about how to grow and sustain support through better stories, "
        "timely thank-yous, thoughtful re-engagement of people who have stepped back, "
        "and creative ideas that feel natural to a ministry. I would be glad to learn "
        "GOD TV’s testimonies, systems and CRM thoroughly, including Salesforce, and "
        "to travel when staff gatherings or events require it. GOD TV’s thirty-year "
        "legacy, and this new season of growth, is something I would be honoured to "
        "serve. I believe my faith, church media experience and commitment to "
        "prayerful, relationship-led communication would make me a strong addition to "
        "the UK Partnerships team.",
        **body_kwargs,
    )

    add_paragraph(
        doc,
        "Thank you for considering my application. I have attached my CV and would "
        "welcome the chance to discuss how I can help partners feel known, encouraged "
        "and invited into what God is doing through this ministry.",
        **body_kwargs,
    )

    add_paragraph(doc, "Yours faithfully,", size=11, space_after=20, line_spacing=1.0)
    add_paragraph(doc, "Dennis Kariuki", size=13, bold=True, color=NAVY, space_after=2, line_spacing=1.0)
    add_paragraph(
        doc,
        "Practising Christian  ·  Church Media  ·  Video  ·  Social Media",
        size=10,
        italic=True,
        color=MUTED,
        space_after=0,
        line_spacing=1.0,
    )

    output = "/workspace/Dennis_Kariuki_GOD_TV_Cover_Letter.docx"
    doc.save(output)
    print(output)


if __name__ == "__main__":
    build()
